
import os
import base64
import streamlit as st
from dotenv import load_dotenv
from openai import OpenAI

from pypdf import PdfReader
from docx import Document
from pptx import Presentation
import pandas as pd


# ==================================================
# 1. PAGE CONFIGURATION
# ==================================================

st.set_page_config(
    page_title="My AI Assistant",
    page_icon="🤖",
    layout="wide"
)


# ==================================================
# 2. CUSTOM CSS
# ==================================================

st.markdown("""
<style>
.stApp {
    background-color: #0e1117;
    color: #ffffff;
}

.main-title {
    text-align: center;
    font-size: 38px;
    font-weight: bold;
    color: #60a5fa;
    margin-top: 10px;
}

.subtitle {
    text-align: center;
    color: #aaaaaa;
    margin-bottom: 25px;
}

section[data-testid="stSidebar"] {
    background-color: #161b22;
}

.stChatMessage {
    border-radius: 12px;
}
</style>
""", unsafe_allow_html=True)


# ==================================================
# 3. LOAD API KEY
# ==================================================

load_dotenv()

api_key = os.getenv("GROQ_API_KEY")

if not api_key:
    st.error(
        "GROQ_API_KEY not found. "
        "Please check your .env file."
    )
    st.stop()


# ==================================================
# 4. CONNECT TO GROQ API
# ==================================================

client = OpenAI(
    api_key=api_key,
    base_url="https://api.groq.com/openai/v1"
)


# ==================================================
# 5. AVAILABLE MODELS
# ==================================================

MODELS = {
    "GPT OSS 120B - Advanced Research": "openai/gpt-oss-120b",
    "GPT OSS 20B - Fast Chat": "openai/gpt-oss-20b",
    "Qwen 3.8 27B - Image Analysis": "qwen/qwen3.8-27b"
}

VISION_MODEL = "qwen/qwen3.8-27b"

WEB_SEARCH_MODELS = [
    "openai/gpt-oss-120b",
    "openai/gpt-oss-20b"
]


# ==================================================
# 6. SESSION STATE
# ==================================================

if "messages" not in st.session_state:
    st.session_state.messages = []

if "document_text" not in st.session_state:
    st.session_state.document_text = ""

if "uploaded_file_names" not in st.session_state:
    st.session_state.uploaded_file_names = []


# ==================================================
# 7. DOCUMENT EXTRACTION FUNCTIONS
# ==================================================

def extract_pdf(file):

    reader = PdfReader(file)

    text = ""

    for page in reader.pages:
        text += (page.extract_text() or "") + "\n"

    return text


def extract_docx(file):

    document = Document(file)

    text = ""

    for paragraph in document.paragraphs:
        text += paragraph.text + "\n"

    return text


def extract_pptx(file):

    presentation = Presentation(file)

    text = ""

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        text += f"\nSlide {slide_number}\n"

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                text += shape.text + "\n"

    return text


def extract_excel(file):

    excel_data = pd.read_excel(
        file,
        sheet_name=None
    )

    text = ""

    for sheet_name, dataframe in excel_data.items():

        text += f"\nSheet: {sheet_name}\n"

        text += dataframe.to_string(
            index=False
        )

    return text


def extract_csv(file):

    dataframe = pd.read_csv(file)

    return dataframe.to_string(
        index=False
    )


def extract_text_file(file):

    return file.getvalue().decode(
        "utf-8",
        errors="replace"
    )


# ==================================================
# 8. PROCESS UPLOADED DOCUMENTS
# ==================================================

def process_uploaded_file(file):

    extension = file.name.rsplit(
        ".",
        1
    )[-1].lower()

    try:

        if extension == "pdf":
            return extract_pdf(file)

        elif extension == "docx":
            return extract_docx(file)

        elif extension == "pptx":
            return extract_pptx(file)

        elif extension == "xlsx":
            return extract_excel(file)

        elif extension == "csv":
            return extract_csv(file)

        elif extension in ["txt", "md", "json"]:
            return extract_text_file(file)

        else:
            st.warning(
                f"Unsupported file type: {file.name}"
            )
            return None

    except Exception as error:

        st.error(
            f"Could not read {file.name}: {error}"
        )

        return None


# ==================================================
# 9. SIDEBAR
# ==================================================

with st.sidebar:

    st.title("⚙️ Settings")

    st.subheader("Choose AI Model")

    selected_model_name = st.selectbox(
        "AI Model",
        list(MODELS.keys())
    )

    selected_model = MODELS[
        selected_model_name
    ]

    st.caption(
        "Use Qwen 3.8 27B when you want to analyse images."
    )

    st.divider()

    # ----------------------------------------------
    # DOCUMENT UPLOAD
    # ----------------------------------------------

    st.subheader("📁 Upload Documents")

    uploaded_files = st.file_uploader(
        "Choose documents",
        type=[
            "pdf",
            "docx",
            "pptx",
            "xlsx",
            "csv",
            "txt",
            "md",
            "json"
        ],
        accept_multiple_files=True,
        key="document_uploader"
    )

    if st.button(
        "Process Documents",
        use_container_width=True
    ):

        if uploaded_files:

            combined_text = ""
            processed_names = []

            with st.spinner(
                "Reading your documents..."
            ):

                for file in uploaded_files:

                    extracted_text = process_uploaded_file(
                        file
                    )

                    if extracted_text:

                        combined_text += (
                            f"\n\nDOCUMENT: {file.name}\n"
                            f"{extracted_text}"
                        )

                        processed_names.append(
                            file.name
                        )

            st.session_state.document_text = (
                combined_text
            )

            st.session_state.uploaded_file_names = (
                processed_names
            )

            if processed_names:
                st.success(
                    f"Processed {len(processed_names)} document(s)."
                )
            else:
                st.warning(
                    "No document text could be extracted."
                )

        else:

            st.warning(
                "Please upload a document first."
            )

    if st.session_state.uploaded_file_names:

        st.write("Processed documents:")

        for name in st.session_state.uploaded_file_names:
            st.caption("📄 " + name)

    st.divider()

    # ----------------------------------------------
    # IMAGE UPLOAD
    # ----------------------------------------------

    st.subheader("🖼️ Upload Image")

    uploaded_image = st.file_uploader(
        "Choose an image",
        type=[
            "png",
            "jpg",
            "jpeg",
            "webp"
        ],
        key="image_uploader"
    )

    if uploaded_image is not None:

        st.image(
            uploaded_image,
            caption="Selected image",
            use_container_width=True
        )

        if selected_model != VISION_MODEL:

            st.warning(
                "Select Qwen 3.8 27B to analyse this image."
            )

    st.divider()

    # ----------------------------------------------
    # CLEAR CHAT
    # ----------------------------------------------

    if st.button(
        "🗑️ Clear Conversation",
        use_container_width=True
    ):

        st.session_state.messages = []

        st.rerun()

    st.divider()

    # ----------------------------------------------
    # DOWNLOAD CHAT
    # ----------------------------------------------

    if st.session_state.messages:

        conversation_text = ""

        for message in st.session_state.messages:

            role = message["role"].upper()

            conversation_text += (
                f"\n{role}:\n"
                f"{message['content']}\n"
            )

        st.download_button(
            label="⬇️ Download Conversation",
            data=conversation_text,
            file_name="chat_history.txt",
            mime="text/plain",
            use_container_width=True
        )

    st.caption("Powered by Python, Streamlit and Groq")


# ==================================================
# 10. MAIN PAGE HEADER
# ==================================================

st.markdown(
    '<div class="main-title">🤖 My AI Assistant</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="subtitle">'
    'Ask questions, analyse images and upload documents'
    '</div>',
    unsafe_allow_html=True
)


# ==================================================
# 11. DISPLAY CHAT HISTORY
# ==================================================

for message in st.session_state.messages:

    with st.chat_message(message["role"]):

        st.markdown(message["content"])


# ==================================================
# 12. CHAT INPUT
# ==================================================

user_input = st.chat_input(
    "Ask me anything..."
)


# ==================================================
# 13. SEND MESSAGE TO AI
# ==================================================

if user_input:

    # Show user message
    with st.chat_message("user"):

        st.markdown(user_input)

    # Save user message as normal text
    st.session_state.messages.append({
        "role": "user",
        "content": user_input
    })

    # ----------------------------------------------
    # SYSTEM INSTRUCTIONS
    # ----------------------------------------------

    system_prompt = """
    You are a helpful AI assistant.

    Answer questions clearly and politely.

    Explain difficult topics in simple language.

    Use examples whenever useful.

    If the user asks about programming,
    provide correct and understandable code.
    """

    # Add extracted document content
    if st.session_state.document_text:

        system_prompt += """

        The user has uploaded documents.

        Use the following document content
        to answer questions about those documents.

        If the answer is not available in the documents,
        clearly tell the user.

        DOCUMENT CONTENT:

        """

        system_prompt += (
            st.session_state.document_text[:30000]
        )

    # ----------------------------------------------
    # PREPARE API MESSAGES
    # ----------------------------------------------

    api_messages = [
        {
            "role": "system",
            "content": system_prompt
        }
    ]

    # Add previous conversation messages.
    # Exclude the latest user message for now.
    api_messages.extend(
        st.session_state.messages[:-1]
    )

    # ----------------------------------------------
    # PREPARE CURRENT MESSAGE
    # ----------------------------------------------

    current_user_message = {
        "role": "user",
        "content": user_input
    }

    # If an image is uploaded and Qwen is selected,
    # send text and image together.
    if uploaded_image is not None:

        if selected_model == VISION_MODEL:

            image_bytes = uploaded_image.getvalue()

            encoded_image = base64.b64encode(
                image_bytes
            ).decode("utf-8")

            image_url = (
                f"data:{uploaded_image.type};base64,"
                f"{encoded_image}"
            )

            current_user_message = {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": user_input
                    },
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": image_url
                        }
                    }
                ]
            }

        else:

            st.warning(
                "Your image was not sent because the selected "
                "model does not support image input. "
                "Choose Qwen 3.8 27B to analyse it."
            )

    # Add the current message to API request
    api_messages.append(
        current_user_message
    )

    # ----------------------------------------------
    # CALL GROQ API
    # ----------------------------------------------

    with st.chat_message("assistant"):

        with st.spinner("AI is thinking..."):

            try:

                response = client.chat.completions.create(
                    model=selected_model,
                    messages=api_messages
                )

                ai_response = (
                    response.choices[0].message.content
                    or "The AI returned an empty response."
                )

                st.markdown(ai_response)

                # Save assistant response
                st.session_state.messages.append({
                    "role": "assistant",
                    "content": ai_response
                })

            except Exception as error:

                st.error(
                    f"Groq API Error: {error}"
                )


# ==================================================
# 14. FOOTER
# ==================================================

st.markdown("---")

st.markdown(
    "<center>Built with Python, Streamlit and Groq 🤖</center>",
    unsafe_allow_html=True
)
